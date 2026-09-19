#!/usr/bin/env python3
"""prerequis.py — ce qui doit être là avant de produire, et ce qui manque.

Le dépôt des skills synchronise le CODE entre les machines. Il ne synchronise
ni les polices, ni les applications, ni les paquets Python : ils vivent hors de
`~/.claude` et chaque machine les installe pour elle-même.

D'où ce script. Sur une machine neuve, ou après une mise à jour de LibreOffice,
il dit en cinq secondes ce qui manque et la commande pour le poser. Sans lui, le
manque se découvre au milieu d'une production, et le plus vicieux ne se découvre
pas du tout : une police absente ne fait pas échouer LibreOffice, elle lui fait
rendre un deck qui n'est pas le vôtre.

    python3 prerequis.py            # verdict
    python3 prerequis.py --corriger # pose ce qui peut se poser sans installer

Code de sortie : 0 si tout est prêt, 1 sinon.
"""

import shutil
import subprocess
import sys
from pathlib import Path

def _trouver_libre():
    """Où est LibreOffice, cherché et non supposé.

    Un chemin codé en dur est un pari sur l'installation de l'autre machine. Il
    a été perdu : Homebrew pose l'application dans `/Applications` ou dans
    `~/Applications` selon les droits, et lie de toute façon un lanceur sur le
    PATH, ce que son propre message d'installation annonce. Une machine où
    `soffice` répondait parfaitement s'est vue déclarer « LibreOffice absent ».

    On demande donc au système, puis on retombe sur les emplacements connus.
    Rend `(bundle, binaire, dossier_de_polices)`, chacun `None` si introuvable.
    """
    # Le BUNDLE d'abord, le lanceur du gestionnaire de paquets ensuite. Les deux
    # savent rendre un PDF, mais seul le bundle porte le dossier de polices où
    # l'on pose Inter. Prendre le lanceur en premier ferait perdre ce dossier et
    # laisserait croire qu'il n'existe pas.
    candidats = [Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
                 Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"]
    lanceur = shutil.which("soffice") or shutil.which("libreoffice")
    if lanceur:
        vrai = Path(lanceur).resolve()
        for parent in vrai.parents:
            if parent.suffix == ".app":
                candidats.insert(0, parent / "Contents/MacOS/soffice")
                break
        candidats.append(vrai)          # dernier recours : le lanceur lui-même

    for c in candidats:
        if not c.exists():
            continue
        # Exister ne suffit pas : il faut RÉPONDRE. Relevé sur une machine où le
        # gestionnaire de paquets avait annoncé l'installation avec succès, posé
        # son lanceur sur le PATH, et jamais l'application. `which` le trouvait,
        # le lanceur échouait à chaque appel sur un binaire absent. Un contrôle
        # qui teste la présence d'un fichier valide une installation cassée.
        try:
            r_ = subprocess.run([str(c), "--version"], capture_output=True,
                                timeout=90)
            if r_.returncode != 0:
                continue
        except OSError:
            continue
        except subprocess.TimeoutExpired:
            pass            # lent n'est pas cassé : le premier lancement l'est
        bundle = next((p for p in c.parents if p.suffix == ".app"), None)
        polices = bundle / "Contents/Resources/fonts/truetype" if bundle else None
        if polices is None or not polices.is_dir():
            # Un lanceur hors bundle : on sait rendre, pas poser de police.
            polices = None
        return bundle, c, polices
    return None, None, None


LIBRE, LIBRE_BIN, LIBRE_FONTS = _trouver_libre()
CHROME = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
if not CHROME.exists():
    _c = shutil.which("google-chrome") or shutil.which("chrome")
    if _c:
        CHROME = Path(_c)

# Les familles que les chaînes de deck posent réellement. Une famille absente
# d'ici n'est pas une erreur : elle n'est simplement pas contrôlée.
FAMILLES = {
    "Inter": (Path.home() / "Library/Fonts", ["Inter-Regular.otf", "Inter-Bold.otf",
                                              "Inter-Italic.otf"]),
    "Georgia": (Path("/System/Library/Fonts/Supplemental"), ["Georgia.ttf",
                                                             "Georgia Italic.ttf"]),
}

OK, KO, MOU = "✅", "⛔", "⚠"


def verifier():
    """Rend une liste de (etat, sujet, detail, remede)."""
    r = []

    # Quel interpréteur, en clair. Une machine peut en porter plusieurs, et un
    # seul avoir les modules : sur une seconde machine, celui du shell de login
    # les avait, celui d'un appel direct non. Sans cette ligne, le verdict est
    # juste pour un Python et muet sur celui qui produira réellement.
    r.append((OK, "interpréteur", f"{sys.executable} ({sys.version.split()[0]})", None))

    for mod, paquet in (("pptx", "python-pptx"), ("PIL", "pillow")):
        try:
            __import__(mod)
            r.append((OK, paquet, "importable", None))
        except ImportError:
            r.append((KO, paquet, "absent", f"pip3 install {paquet}"))

    r.append((OK, "Chrome", "présent", None) if CHROME.exists() else
             (MOU, "Chrome", "absent, le rendu en flux ne marchera pas",
              "brew install --cask google-chrome"))

    # Les polices SYSTÈME d'abord, et indépendamment de LibreOffice. Le moteur de
    # mesure ouvre le fichier de police réel : sans lui, il ne rend pas un
    # résultat approximatif, il lève. Une version antérieure de ce script
    # retournait dès que LibreOffice manquait et ne les regardait jamais : elle
    # annonçait un seul prérequis bloquant quand il pouvait y en avoir deux.
    for famille, (source, fichiers) in FAMILLES.items():
        manque = [f for f in fichiers if not (source / f).exists()]
        if not manque:
            r.append((OK, f"{famille} (système)", f"présente dans {source}", None))
        elif famille == "Inter":
            r.append((KO, f"{famille} (système)",
                      f"{len(manque)} fichier(s) manquant(s) dans {source} — le "
                      f"moteur de mesure LÈVE sans elle",
                      "brew install --cask font-inter"))
        else:
            r.append((MOU, f"{famille} (système)",
                      f"{len(manque)} fichier(s) manquant(s) dans {source}",
                      "police système macOS, normalement présente"))

    if LIBRE_BIN is None:
        trouve = shutil.which("soffice") or shutil.which("libreoffice")
        if trouve:
            r.append((KO, "LibreOffice",
                      f"un lanceur existe en {trouve} mais il NE RÉPOND PAS : "
                      f"l'application n'est pas installée là où il la cherche",
                      "brew reinstall --cask libreoffice"))
        else:
            r.append((KO, "LibreOffice",
                      "absent, aucune vérification visuelle possible",
                      "brew install --cask libreoffice"))
        return r

    r.append((OK, "LibreOffice", f"présent · {LIBRE_BIN}", None))
    if LIBRE_FONTS is None:
        r.append((MOU, "polices LibreOffice",
                  "dossier de polices du bundle introuvable",
                  "le rendu emploiera les polices du système"))
        return r

    # Le piège : LibreOffice ne résout pas les polices utilisateur de macOS. Il
    # ne le signale pas, il substitue. Un deck rendu avec une autre police a
    # d'autres largeurs, donc d'autres retours à la ligne, donc un contrôle
    # visuel qui ne contrôle rien.
    for famille, (source, fichiers) in FAMILLES.items():
        presentes = [f for f in fichiers if (LIBRE_FONTS / f).exists()]
        if len(presentes) == len(fichiers):
            r.append((OK, f"{famille} pour LibreOffice", "installée", None))
            continue
        dispo = [f for f in fichiers if (source / f).exists()]
        if not dispo:
            r.append((KO, f"{famille} pour LibreOffice",
                      f"introuvable même dans {source}",
                      "installer la police sur le système d'abord"))
        else:
            manque = [f for f in fichiers if f not in presentes]
            r.append((MOU, f"{famille} pour LibreOffice",
                      f"{len(manque)} fichier(s) manquant(s) : LibreOffice "
                      f"SUBSTITUERA en silence",
                      f'cp "{source}"/{famille}* "{LIBRE_FONTS}"/'))
    return r


# Les familles dont l'absence bloque, et celles qui dégradent seulement. Inter
# porte tout le corps du texte ; Georgia ne sert qu'aux citations en italique de
# certains profils. Confondre les deux fait échouer une machine pour une police
# décorative.
BLOQUANTES = {"Inter"}


def corriger():
    """Pose ce qui se pose sans rien installer : les polices dans LibreOffice.

    Copie le CONTENU seul, jamais les métadonnées. `shutil.copy2` recopie aussi
    les permissions et les attributs étendus, et il échoue sur un fichier du
    volume système signé — ce qui est exactement le cas des polices macOS
    livrées avec l'OS. Relevé sur une seconde machine : Inter, qui vient du
    dossier utilisateur, passait ; Georgia, qui vient de /System, échouait.

    Et une famille qui échoue n'arrête pas les autres : la version précédente
    rendait la main au premier refus, donc sans verdict, alors que l'essentiel
    était peut-être déjà posé.
    """
    if LIBRE_BIN is None:
        print(f"  {KO} LibreOffice absent, rien à corriger ici")
        return 1
    if LIBRE_FONTS is None:
        print(f"  {MOU} LibreOffice trouvé en {LIBRE_BIN}, mais sans dossier de\n       polices : rien à poser, le rendu emploiera les polices du système")
        return 0
    pose, echecs = 0, []
    for famille, (source, fichiers) in FAMILLES.items():
        for f in fichiers:
            src, dst = source / f, LIBRE_FONTS / f
            if not src.exists() or dst.exists():
                continue
            try:
                shutil.copyfile(src, dst)
                print(f"  {OK} {f} → LibreOffice")
                pose += 1
            except OSError as e:
                marque = KO if famille in BLOQUANTES else MOU
                print(f"  {marque} {f} : {type(e).__name__}")
                echecs.append((famille, src, marque))

    print(f"\n  {pose} police(s) posée(s)" if pose else "\n  rien à poser")
    for famille, src, marque in echecs:
        if marque == MOU:
            print(f"  {MOU} {famille} n'a pas pu être posée. Elle ne porte que "
                  f"des citations : le rendu reste fidèle sur le corps du texte.")
        print(f"       à la main : sudo cp \"{src}\" \"{LIBRE_FONTS}\"/")
    return 1 if any(m == KO for _, _, m in echecs) else 0


def preuve_rendu():
    """Rend un PPTX d'une slide et lit les polices réellement embarquées.

    C'est la seule vérification qui prouve quelque chose. Le reste dit que les
    fichiers sont là ; celle-ci dit que le rendu les emploie.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return None
    import re
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "sonde.pptx"
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.3333), Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        b = s.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1))
        run = b.text_frame.paragraphs[0].add_run()
        run.text = "Sonde de police"
        run.font.name, run.font.size = "Inter", Pt(28)
        prs.save(str(p))
        try:
            # Large : le PREMIER lancement de LibreOffice construit son profil
            # utilisateur et prend bien plus longtemps que les suivants. Un
            # timeout calibré sur le régime établi fait échouer la machine neuve,
            # c'est-à-dire exactement celle qu'on vérifie.
            subprocess.run([str(LIBRE_BIN), "--headless", "--convert-to", "pdf",
                            "--outdir", d, str(p)],
                           check=True, capture_output=True, timeout=300)
        except subprocess.TimeoutExpired:
            return ("échec", "LibreOffice n'a pas répondu en 5 minutes. Le "
                    "lancer une fois à la main puis relancer ce script.")
        except subprocess.CalledProcessError as e:
            fin = (e.stderr or b"").decode(errors="replace").strip()[-160:]
            return ("échec", f"LibreOffice a refusé la conversion. {fin}")
        except OSError as e:
            return ("échec", f"{type(e).__name__}: {e}")
        pdf = Path(d) / "sonde.pdf"
        if not pdf.exists():
            return ("échec", "aucun PDF produit, sans message d'erreur. Lancer "
                    "LibreOffice une fois à la main puis relancer ce script.")
        familles = {m.decode().split("+")[-1].split("-")[0] for m in
                    re.findall(rb"/BaseFont\s*/([A-Za-z0-9+\-,_]+)", pdf.read_bytes())}
        return ("ok", familles)


def main():
    if "--corriger" in sys.argv:
        print("\n  POSE DES POLICES DANS LIBREOFFICE")
        code = corriger()
        print()
        if code:
            return code

    resultats = verifier()
    print("\n  PRÉREQUIS DE PRODUCTION D'UN DECK\n")
    for etat, sujet, detail, remede in resultats:
        print(f"  {etat} {sujet:<28} {detail}")
        if remede:
            print(f"       → {remede}")

    if LIBRE_BIN is not None:
        etat, detail = preuve_rendu()
        print()
        if etat == "échec":
            # Elle ne se laisse pas ignorer. C'est la SEULE vérification qui
            # prouve quelque chose : les autres disent que les fichiers sont là,
            # celle-ci dit que le rendu les emploie. Conclure « tout est prêt »
            # sans elle serait un faux vert, et un faux vert vaut moins que rien.
            print(f"  {KO} preuve de rendu             non établie · {detail}")
            resultats.append((KO, "preuve de rendu", detail, None))
        elif "Inter" in detail:
            print(f"  {OK} preuve de rendu             Inter est bien tracée "
                  f"(polices vues : {', '.join(sorted(detail))})")
        else:
            print(f"  {KO} preuve de rendu             Inter est SUBSTITUÉE par "
                  f"{', '.join(sorted(detail))}")
            print(f"       → python3 {Path(__file__).name} --corriger")
            resultats.append((KO, "rendu", "substitution", None))

    bloquants = [x for x in resultats if x[0] == KO]
    print()
    if bloquants:
        print(f"  {KO} {len(bloquants)} prérequis bloquant(s). "
              f"Produire un deck maintenant donnerait un résultat faux.")
        return 1
    mous = [x for x in resultats if x[0] == MOU]
    if mous:
        print(f"  {MOU} tout l'essentiel est là, {len(mous)} point(s) à surveiller.")
    else:
        print(f"  {OK} tout est prêt.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
